"""Tests that the validator catches the failures it exists to catch.

A validator that always passes is worse than none, so each test breaks the
workbook in one specific way and asserts the report notices.
"""
from __future__ import annotations

import json
import os
import shutil
import tempfile
import unittest

import openpyxl

from .client import CallRecord
from .validate import ERROR, to_json, validate
from .workbook import applicable_sheets, map_sheet, state_row, write_results

HERE = os.path.dirname(__file__)
TEMPLATE = os.path.join(
    HERE, "..", "..", "evidence", "metrc",
    "Generic_Evaluation_for_All_States_MASTER_10.2025.xlsx",
)


def record(sheet, step, **kw):
    base = dict(
        sheet=sheet, step=step, method="POST", path="/locations/v2/",
        url="https://sandbox-api-ny.metrc.com/locations/v2/?licenseNumber=X-1",
        status=200, license_number="X-1", request_body=[{"Name": "L"}],
        response_body={"Ids": [7]}, object_ids=[7], names=["L"],
    )
    base.update(kw)
    return CallRecord(**base)


class _Fixture(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.out = os.path.join(self.dir, "filled.xlsx")
        self.records = [
            record("Locations", "Step 1"),
            record("Locations", "Step 2", method="PUT"),
            record("Locations", "Step 3", method="GET", request_body=None),
        ]

        class Fake:
            def __init__(self, records):
                self.records = records

            def for_step(self, sheet, step):
                hits = [r for r in self.records if r.sheet == sheet and r.step == step]
                return hits[-1] if hits else None

        write_results(TEMPLATE, self.out, Fake(self.records), only_sheets={"Locations"})

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def report(self):
        return validate(self.out, self.records, sheets=["Locations"], template_path=TEMPLATE)


class TestValidate(_Fixture):
    def test_a_correctly_filled_sheet_passes(self):
        report = self.report()
        self.assertTrue(report.ok(), [str(f) for f in report.errors])
        self.assertEqual(report.steps_ok, 3)

    def test_a_tampered_cell_is_caught(self):
        wb = openpyxl.load_workbook(self.out)
        ws = wb["Locations"]
        col = map_sheet(ws).columns["status"]
        ws[f"{col}5"] = "999"
        wb.save(self.out)
        report = self.report()
        self.assertTrue(any("transcript" in f.message for f in report.errors), report.findings)

    def test_a_cleared_cell_is_caught(self):
        wb = openpyxl.load_workbook(self.out)
        ws = wb["Locations"]
        col = map_sheet(ws).columns["evidence"]
        ws[f"{col}5"] = None
        wb.save(self.out)
        self.assertFalse(self.report().ok())

    def test_unminified_evidence_is_caught(self):
        wb = openpyxl.load_workbook(self.out)
        ws = wb["Locations"]
        col = map_sheet(ws).columns["evidence"]
        ws[f"{col}5"] = '[\n  {"Name": "L"}\n]'
        wb.save(self.out)
        messages = " ".join(f.message for f in self.report().findings if f.level == ERROR)
        self.assertIn("transcript", messages)

    def test_a_non_https_url_is_caught(self):
        broken = [record("Locations", "Step 1", url="http://insecure/locations/v2/")]
        broken += self.records[1:]
        wb = openpyxl.load_workbook(self.out)
        ws = wb["Locations"]
        ws[f"{map_sheet(ws).columns['url']}5"] = "http://insecure/locations/v2/"
        wb.save(self.out)
        report = validate(self.out, broken, sheets=["Locations"], template_path=TEMPLATE)
        self.assertTrue(any("not https" in f.message for f in report.errors), report.findings)

    def test_json_export_mirrors_the_workbook(self):
        payload = to_json(self.out, self.records, sheets=["Locations"])
        self.assertEqual(payload["summary"]["steps_total"], 3)
        self.assertEqual(payload["summary"]["steps_ok"], 3)
        steps = payload["sheets"][0]["steps"]
        self.assertEqual([s["step"] for s in steps], ["Step 1", "Step 2", "Step 3"])
        self.assertEqual(steps[0]["object_ids"], [7])
        self.assertTrue(steps[0]["task"])
        json.dumps(payload)  # must be serialisable


class TestStateScope(unittest.TestCase):
    def test_ny_is_open_loop_and_uses_the_patient_sales_tab(self):
        sheets = applicable_sheets(TEMPLATE, "NY")
        self.assertIn("PlantBatches", sheets)
        self.assertNotIn("Closed Loop States PlantBatches", sheets)
        self.assertIn("Sales with Patient Look Up", sheets)
        self.assertIn("LabResults", sheets)
        self.assertNotIn("CA ONLY Labs", sheets)
        self.assertNotIn("CA- SalesRetailDeliveries", sheets)

    def test_ca_is_closed_loop_and_uses_the_ca_only_tabs(self):
        sheets = applicable_sheets(TEMPLATE, "CA")
        self.assertIn("Closed Loop States PlantBatches", sheets)
        self.assertIn("CA ONLY Labs", sheets)
        self.assertIn("CA- SalesRetailDeliveries", sheets)
        self.assertNotIn("PlantBatches", sheets)

    def test_an_unlisted_state_is_rejected_rather_than_guessed(self):
        with self.assertRaises(KeyError):
            state_row(TEMPLATE, "ZZ")


if __name__ == "__main__":
    unittest.main()


class TestPermissions(unittest.TestCase):
    """The Permissions tab decides production access, so a request that omits
    part of its own dependency chain must not pass silently."""

    def setUp(self):
        self.wb = openpyxl.load_workbook(TEMPLATE)

    def test_a_complete_sales_request_reports_no_gaps(self):
        from .workbook import fill_permissions
        request = {
            "facility_types": ["Sales"],
            "areas": {a: {"get": True, "write": False} for a in [
                "Strains", "Items", "Packages", "Sales", "Sales Deliveries",
                "GET Transfers / Wholesale",
            ]},
        }
        self.assertEqual(fill_permissions(self.wb, request), [])

    def test_an_incomplete_request_names_every_missing_dependency(self):
        from .workbook import fill_permissions
        gaps = fill_permissions(
            self.wb, {"facility_types": ["Sales"], "areas": {"Sales": {"get": True}}}
        )
        joined = " ".join(gaps)
        for required in ("Strains", "Packages", "Items", "GET Transfers"):
            self.assertIn(required, joined)

    def test_marks_land_in_the_get_and_write_columns(self):
        from .workbook import fill_permissions, _permission_rows
        fill_permissions(self.wb, {
            "facility_types": [],
            "areas": {"Packages": {"get": True, "write": True},
                      "Strains": {"get": True, "write": False}},
        })
        ws = self.wb["Permissions"]
        rows = _permission_rows(ws)
        self.assertEqual(ws.cell(rows["Packages"], 3).value, "X")
        self.assertEqual(ws.cell(rows["Packages"], 4).value, "X")
        self.assertEqual(ws.cell(rows["Strains"], 3).value, "X")
        self.assertIsNone(ws.cell(rows["Strains"], 4).value)
